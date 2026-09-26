from pathlib import Path
from textwrap import dedent
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path.cwd() / "リネームAPP"

RENAME_PY = r'''from __future__ import annotations

import csv
import re
import shutil
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

APP_ROOT = Path(__file__).resolve().parents[1]
CONFIG_FILE = APP_ROOT / "config" / "RenameConfig.xlsx"
TARGET_DIR = APP_ROOT / "target"
OUTPUT_DIR = APP_ROOT / "output"
LOG_DIR = APP_ROOT / "logs"

TRUE_VALUES = {"true", "1", "yes", "on", "はい", "有効"}
INVALID_WINDOWS_CHARS = '<>:"/\\|?*'


def as_bool(value) -> bool:
    return str(value).strip().lower() in TRUE_VALUES


def load_config() -> dict:
    if not CONFIG_FILE.exists():
        raise FileNotFoundError(f"設定ファイルがありません: {CONFIG_FILE}")
    wb = load_workbook(CONFIG_FILE, data_only=True)
    if "設定" not in wb.sheetnames:
        raise ValueError("RenameConfig.xlsx に『設定』シートがありません。")
    ws = wb["設定"]
    config = {}
    for key, value, *_ in ws.iter_rows(min_row=2, values_only=True):
        if key:
            config[str(key).strip()] = value
    return config


def safe_text(value) -> str:
    return "" if value is None else str(value)


def normalize_extension(ext: str) -> str:
    ext = ext.strip()
    if not ext or ext == "*":
        return ""
    return ext if ext.startswith(".") else "." + ext


def allowed_extension(path: Path, setting: str) -> bool:
    if not setting or setting.strip() == "*":
        return True
    allowed = {normalize_extension(x).lower() for x in setting.split(",") if x.strip()}
    return path.suffix.lower() in allowed


def sanitize_filename(name: str, replacement: str) -> str:
    for char in INVALID_WINDOWS_CHARS:
        name = name.replace(char, replacement)
    return name.rstrip(". ")


def apply_case(text: str, mode: str) -> str:
    mode = mode.strip().lower()
    if mode in {"lower", "小文字"}:
        return text.lower()
    if mode in {"upper", "大文字"}:
        return text.upper()
    return text


@dataclass
class Plan:
    source: Path
    destination: Path
    status: str
    reason: str = ""


def build_new_name(path: Path, config: dict, index: int) -> str:
    stem = path.stem
    suffix = path.suffix

    before = safe_text(config.get("置換前"))
    after = safe_text(config.get("置換後"))
    use_regex = as_bool(config.get("正規表現を使用", False))
    if before:
        stem = re.sub(before, after, stem) if use_regex else stem.replace(before, after)

    if as_bool(config.get("空白を削除", False)):
        stem = re.sub(r"\s+", "", stem)

    stem = safe_text(config.get("先頭文字")) + stem + safe_text(config.get("末尾文字"))

    if as_bool(config.get("日付追加", False)):
        date_format = safe_text(config.get("日付形式")) or "%Y%m%d"
        date_text = datetime.now().strftime(date_format)
        position = safe_text(config.get("日付位置")) or "先頭"
        stem = f"{date_text}_{stem}" if position == "先頭" else f"{stem}_{date_text}"

    if as_bool(config.get("連番追加", False)):
        start = int(config.get("連番開始", 1) or 1)
        digits = int(config.get("連番桁数", 3) or 3)
        number = str(start + index).zfill(digits)
        position = safe_text(config.get("連番位置")) or "末尾"
        stem = f"{number}_{stem}" if position == "先頭" else f"{stem}_{number}"

    stem = apply_case(stem, safe_text(config.get("英字大小変換")))
    invalid_replacement = safe_text(config.get("禁止文字の置換文字")) or "_"
    new_name = sanitize_filename(stem + suffix, invalid_replacement)
    return new_name


def collect_files(config: dict) -> list[Path]:
    recursive = as_bool(config.get("サブフォルダを含む", False))
    include_hidden = as_bool(config.get("隠しファイルを含む", False))
    extension_setting = safe_text(config.get("対象拡張子")) or "*"
    iterator = TARGET_DIR.rglob("*") if recursive else TARGET_DIR.iterdir()
    files = []
    for path in iterator:
        if not path.is_file():
            continue
        if not include_hidden and any(part.startswith(".") for part in path.relative_to(TARGET_DIR).parts):
            continue
        if allowed_extension(path, extension_setting):
            files.append(path)
    return sorted(files, key=lambda p: str(p).lower())


def build_plans(config: dict) -> list[Plan]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    files = collect_files(config)
    plans = []
    destinations = set()
    overwrite = as_bool(config.get("同名ファイルを上書き", False))

    for index, source in enumerate(files):
        new_name = build_new_name(source, config, index)
        destination = OUTPUT_DIR / new_name
        status, reason = "OK", ""

        if not new_name or new_name == source.name and source.parent == OUTPUT_DIR:
            status, reason = "SKIP", "変更なし"
        elif destination in destinations:
            status, reason = "ERROR", "変換後の名前が重複"
        elif destination.exists() and not overwrite:
            status, reason = "ERROR", "出力先に同名ファイルが存在"

        destinations.add(destination)
        plans.append(Plan(source, destination, status, reason))
    return plans


def write_preview(plans: list[Plan]) -> Path:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    preview_file = LOG_DIR / "preview.csv"
    with preview_file.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["状態", "変更前", "変更後", "理由"])
        for p in plans:
            writer.writerow([p.status, str(p.source), str(p.destination), p.reason])
    return preview_file


def show_preview(plans: list[Plan]) -> None:
    print("\n===== 安全プレビュー（この時点では変更しません） =====")
    if not plans:
        print("対象ファイルがありません。target フォルダと対象拡張子を確認してください。")
        return
    for p in plans:
        print(f"[{p.status}] {p.source.name} -> {p.destination.name}" + (f" ({p.reason})" if p.reason else ""))
    ok = sum(p.status == "OK" for p in plans)
    errors = sum(p.status == "ERROR" for p in plans)
    print(f"\n対象: {len(plans)}件 / 実行可能: {ok}件 / エラー: {errors}件")


def execute(plans: list[Plan], config: dict) -> Path:
    if any(p.status == "ERROR" for p in plans):
        raise RuntimeError("エラー項目があるため実行を中止しました。preview.csv を確認してください。")

    mode = safe_text(config.get("処理方式")) or "コピー"
    overwrite = as_bool(config.get("同名ファイルを上書き", False))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_file = LOG_DIR / f"result_{timestamp}.csv"

    with log_file.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["結果", "変更前", "変更後", "処理方式"])
        for p in plans:
            if p.status != "OK":
                continue
            if p.destination.exists() and overwrite:
                p.destination.unlink()
            if mode == "移動":
                shutil.move(str(p.source), str(p.destination))
            else:
                shutil.copy2(p.source, p.destination)
            writer.writerow(["成功", str(p.source), str(p.destination), mode])
    return log_file


def main() -> int:
    try:
        TARGET_DIR.mkdir(parents=True, exist_ok=True)
        config = load_config()
        plans = build_plans(config)
        preview_file = write_preview(plans)
        show_preview(plans)
        print(f"プレビューCSV: {preview_file}")

        if not plans or not any(p.status == "OK" for p in plans):
            return 0
        if not as_bool(config.get("実行を許可", False)):
            print("\n設定の『実行を許可』が FALSE のため、プレビューのみで終了します。")
            return 0
        if any(p.status == "ERROR" for p in plans):
            print("\nエラーがあるため実行できません。")
            return 1

        confirm = input("\n内容を確認しましたか？ 実行する場合だけ EXECUTE と入力: ").strip()
        if confirm != "EXECUTE":
            print("キャンセルしました。ファイルは変更されていません。")
            return 0

        log_file = execute(plans, config)
        print(f"完了しました。結果ログ: {log_file}")
        return 0
    except Exception as exc:
        print(f"エラー: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
'''

README = '''# リネームAPP

Excelの設定を読み込み、Pythonで安全にファイル名を変換する試作版です。

## 初回手順

1. `target` にテスト用ファイルを入れます。
2. `config/RenameConfig.xlsx` の青文字セルを変更します。
3. 最初は `実行を許可` を `FALSE` のままにします。
4. `run_preview.bat` を実行し、`logs/preview.csv` を確認します。
5. 問題がなければ `実行を許可` を `TRUE` にし、再実行します。
6. 実行確認では `EXECUTE` と完全一致で入力した場合のみ処理します。

## 安全設計

- 初期値はプレビュー専用です。
- 初期の処理方式は「コピー」で、元ファイルを残します。
- 同名ファイルの上書きは初期状態で禁止です。
- 重複名や衝突がある場合は処理を中止します。
- 実行前後の内容をCSVログへ保存します。
'''

RUN_BAT = r'''@echo off
chcp 65001 > nul
cd /d "%~dp0"
python app\rename.py
pause
'''


def create_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "設定"
    ws.sheet_view.showGridLines = False
    headers = ["設定項目", "設定値", "説明", "入力例"]
    ws.append(headers)
    rows = [
        ("実行を許可", False, "FALSEならプレビューのみ。TRUEでも最後にEXECUTE入力が必要", "FALSE"),
        ("処理方式", "コピー", "コピー推奨。移動は元ファイルが移動します", "コピー / 移動"),
        ("置換前", "IMG_", "検索する文字。空欄なら置換しません", "IMG_"),
        ("置換後", "写真_", "置き換える文字", "写真_"),
        ("正規表現を使用", False, "初心者はFALSE推奨", "FALSE"),
        ("先頭文字", "", "ファイル名の先頭に追加", "旅行_"),
        ("末尾文字", "", "拡張子の直前に追加", "_確認済"),
        ("空白を削除", False, "半角・全角を含む空白文字を削除", "FALSE"),
        ("日付追加", False, "現在日付を追加", "FALSE"),
        ("日付形式", "%Y%m%d", "Pythonのstrftime形式", "%Y%m%d"),
        ("日付位置", "先頭", "先頭または末尾", "先頭"),
        ("連番追加", False, "並び順に連番を追加", "FALSE"),
        ("連番開始", 1, "連番の開始番号", "1"),
        ("連番桁数", 3, "001のように揃える桁数", "3"),
        ("連番位置", "末尾", "先頭または末尾", "末尾"),
        ("英字大小変換", "変更なし", "変更なし・小文字・大文字", "変更なし"),
        ("対象拡張子", "*", "* は全て。複数はカンマ区切り", ".jpg,.png"),
        ("サブフォルダを含む", False, "target配下を再帰検索", "FALSE"),
        ("隠しファイルを含む", False, "通常はFALSE推奨", "FALSE"),
        ("同名ファイルを上書き", False, "安全のためFALSE推奨", "FALSE"),
        ("禁止文字の置換文字", "_", "Windowsで使えない文字を置き換える文字", "_"),
    ]
    for row in rows:
        ws.append(row)

    dark = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="DDEBF7")
    caution_fill = PatternFill("solid", fgColor="FCE4D6")
    for cell in ws[1]:
        cell.fill = dark
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).fill = input_fill
        ws.cell(row, 2).font = Font(color="0000FF")
    ws["B2"].fill = caution_fill
    ws["B2"].font = Font(color="C00000", bold=True)
    ws.column_dimensions["A"].width = 23
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["D"].width = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"

    bool_dv = DataValidation(type="list", formula1='"TRUE,FALSE"')
    ws.add_data_validation(bool_dv)
    for r in [2, 6, 9, 10, 13, 19, 20, 21]:
        bool_dv.add(ws.cell(r, 2))
    mode_dv = DataValidation(type="list", formula1='"コピー,移動"')
    ws.add_data_validation(mode_dv); mode_dv.add(ws["B3"])
    pos_dv = DataValidation(type="list", formula1='"先頭,末尾"')
    ws.add_data_validation(pos_dv); pos_dv.add(ws["B12"]); pos_dv.add(ws["B16"])
    case_dv = DataValidation(type="list", formula1='"変更なし,小文字,大文字"')
    ws.add_data_validation(case_dv); case_dv.add(ws["B17"])

    guide = wb.create_sheet("使い方")
    guide.sheet_view.showGridLines = False
    guide["A1"] = "安全な使い方"
    guide["A1"].font = Font(size=16, bold=True, color="1F4E78")
    steps = [
        "1. target フォルダへコピーしたテストファイルを入れる",
        "2. 設定シートの青文字セルを編集する",
        "3. 実行を許可=FALSE のままプレビューする",
        "4. logs/preview.csv で変更前と変更後を確認する",
        "5. 問題がなければ実行を許可=TRUEにする",
        "6. 実行時に EXECUTE と入力する",
        "7. output と result_日時.csv を確認する",
    ]
    for i, text in enumerate(steps, start=3):
        guide.cell(i, 1, text)
    guide.column_dimensions["A"].width = 72
    wb.save(path)


def main():
    for rel in ["app", "config", "ui", "logs", "target", "output"]:
        (ROOT / rel).mkdir(parents=True, exist_ok=True)
    (ROOT / "app" / "rename.py").write_text(RENAME_PY, encoding="utf-8")
    (ROOT / "README.md").write_text(README, encoding="utf-8")
    (ROOT / "run_preview.bat").write_text(RUN_BAT, encoding="utf-8")
    (ROOT / "ui" / "ここにRenameUI.pptmを作成.txt").write_text(
        "PowerPoint UIは次段階で作成します。まずExcel + Pythonの安全動作を確認してください。\n",
        encoding="utf-8",
    )
    for rel in ["logs", "target", "output"]:
        (ROOT / rel / ".gitkeep").write_text("", encoding="utf-8")
    create_workbook(ROOT / "config" / "RenameConfig.xlsx")
    print(f"作成しました: {ROOT}")


if __name__ == "__main__":
    main()
